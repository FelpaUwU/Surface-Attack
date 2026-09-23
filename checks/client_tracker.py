"""Marks generated reports back into the sales team's client-tracking
spreadsheet: column "Informe" gets "SI" and column "Scoring" gets the
report's grade letter, matched by domain.

A timestamped backup of the workbook is made before every write, since
saving through openpyxl can drop extensions (data validation, conditional
formatting) that other sheets in a hand-built workbook may rely on.
"""
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl


class TrackerLayoutError(ValueError):
    """The expected 'dominio' / 'Informe' / 'Scoring' columns weren't found."""


def _find_columns(ws):
    for row in ws.iter_rows(min_row=1, max_row=15):
        headers = {cell.value: cell.column for cell in row if isinstance(cell.value, str)}
        if "dominio" in headers:
            return row[0].row, headers.get("dominio"), headers.get("Informe"), headers.get("Scoring")
    return None, None, None, None


def update_tracker(xlsx_path, sheet_name, domain_grades):
    """domain_grades: iterable of (domain, grade) pairs for successfully generated reports.

    Returns (matched, total, backup_path).
    """
    xlsx_path = Path(xlsx_path)
    backup_path = xlsx_path.with_name(f"{xlsx_path.stem}_backup_{datetime.now():%Y%m%d_%H%M%S}{xlsx_path.suffix}")
    shutil.copy2(xlsx_path, backup_path)

    wb = openpyxl.load_workbook(xlsx_path)
    if sheet_name not in wb.sheetnames:
        raise TrackerLayoutError(f"La hoja '{sheet_name}' no existe en {xlsx_path.name}.")
    ws = wb[sheet_name]

    header_row, domain_col, informe_col, scoring_col = _find_columns(ws)
    if not (header_row and domain_col and informe_col and scoring_col):
        raise TrackerLayoutError(
            f"No se encontraron las columnas 'dominio' / 'Informe' / 'Scoring' en la hoja '{sheet_name}'."
        )

    by_domain = {domain.strip().lower(): grade for domain, grade in domain_grades if domain}

    matched = 0
    for row in ws.iter_rows(min_row=header_row + 1):
        cell_domain = row[domain_col - 1].value
        if not cell_domain:
            continue
        key = str(cell_domain).strip().lower()
        if key in by_domain:
            row[informe_col - 1].value = "SI"
            row[scoring_col - 1].value = by_domain[key]
            matched += 1

    wb.save(xlsx_path)
    return matched, len(by_domain), backup_path
